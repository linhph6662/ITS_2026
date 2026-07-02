import csv
import os
from typing import List, Dict, Any
from utils.logger import logger

def export_to_csv(data: List[Dict[str, Any]], file_path: str) -> bool:
    """
    Exports vehicle history records to a CSV file.
    
    Args:
        data: List of dictionary entries from the database
        file_path: Save destination path
        
    Returns:
        True if export succeeded, False otherwise.
    """
    try:
        logger.info("Starting CSV export to destination path: %s", file_path)
        os.makedirs(os.path.dirname(os.path.abspath(file_path)), exist_ok=True)
        headers = ["Vehicle ID", "Image Path", "Speed (km/h)", "DateTime", "Video Name"]
        
        with open(file_path, 'w', newline='', encoding='utf-8') as f:
            writer = csv.writer(f)
            writer.writerow(headers)
            for row in data:
                writer.writerow([
                    row.get("vehicle_id", "N/A"),
                    row.get("image_path", "N/A"),
                    f"{row.get('speed', 0.0):.2f}",
                    row.get("datetime", "N/A"),
                    row.get("video_name", "N/A")
                ])
        logger.info("Successfully exported %d records to CSV.", len(data))
        return True
    except Exception as e:
        logger.error("Failed to compile CSV report: %s", e)
        return False

def export_to_excel(data: List[Dict[str, Any]], file_path: str) -> bool:
    """
    Exports vehicle history records to an Excel file (.xlsx) using openpyxl.
    Falls back to CSV format writing on failure.
    
    Args:
        data: List of dictionary entries
        file_path: Target Excel file destination path
        
    Returns:
        True if success, False otherwise.
    """
    try:
        logger.info("Starting Excel export to path: %s", file_path)
        os.makedirs(os.path.dirname(os.path.abspath(file_path)), exist_ok=True)
        import openpyxl
        from openpyxl.styles import Font, Alignment, PatternFill
        
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Vehicle Detection Logs"
        
        headers = ["Vehicle ID", "Image Path", "Speed (km/h)", "DateTime", "Video Name"]
        ws.append(headers)
        
        # Cyberpunk/Neon header styles matching GUI theme (#00ADB5)
        header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="00ADB5", end_color="00ADB5", fill_type="solid")
        
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")
            
        # Add records
        for row in data:
            ws.append([
                row.get("vehicle_id", "N/A"),
                row.get("image_path", "N/A"),
                row.get("speed", 0.0),
                row.get("datetime", "N/A"),
                row.get("video_name", "N/A")
            ])
            
        # Auto-adjust column widths dynamically to prevent cell text clipping
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        wb.save(file_path)
        logger.info("Successfully exported %d records to Excel.", len(data))
        return True
    except Exception as e:
        logger.warning("openpyxl Excel export crashed (%s). Falling back to standard CSV format.", e)
        csv_path = file_path.rsplit('.', 1)[0] + '.csv'
        return export_to_csv(data, csv_path)

def export_to_pdf(data: List[Dict[str, Any]], file_path: str) -> bool:
    """
    Exports vehicle history records to a styled PDF file using ReportLab flowables.
    
    Args:
        data: List of database dictionary records
        file_path: PDF output file destination path
        
    Returns:
        True if success, False otherwise.
    """
    try:
        logger.info("Starting PDF report generation to path: %s", file_path)
        os.makedirs(os.path.dirname(os.path.abspath(file_path)), exist_ok=True)
        from reportlab.lib.pagesizes import letter
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib import colors
        
        doc = SimpleDocTemplate(
            file_path, 
            pagesize=letter, 
            leftMargin=30, 
            rightMargin=30, 
            topMargin=30, 
            bottomMargin=30
        )
        story = []
        
        # Cyber theme styled title and subtitle
        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            'TitleStyle',
            parent=styles['Heading1'],
            fontName='Helvetica-Bold',
            fontSize=20,
            textColor=colors.HexColor('#00ADB5'),
            spaceAfter=5
        )
        
        sub_style = ParagraphStyle(
            'SubStyle',
            parent=styles['Normal'],
            fontName='Helvetica-Oblique',
            fontSize=10,
            textColor=colors.HexColor('#666666'),
            spaceAfter=15
        )
        
        story.append(Paragraph("Traffic Monitoring System Logs", title_style))
        story.append(Paragraph("Automated Vehicle Tracking and Speed Estimation Report", sub_style))
        story.append(Spacer(1, 10))
        
        # Header columns
        headers = ["Vehicle ID", "Image Path", "Speed", "DateTime", "Video Source"]
        table_data = [headers]
        
        for row in data:
            img_path = row.get("image_path", "N/A")
            # Truncate overly long absolute directories
            if img_path and len(img_path) > 30:
                img_path = "..." + img_path[-27:]
                
            table_data.append([
                str(row.get("vehicle_id", "N/A")),
                img_path,
                f"{row.get('speed', 0.0):.1f} km/h",
                str(row.get("datetime", "N/A")),
                str(row.get("video_name", "N/A"))
            ])
            
        # Standard letter width = 612. Printable grid width = 552.
        # Grid sizes: ID (60), Image Path (150), Speed (80), DateTime (142), Video (120)
        t = Table(table_data, colWidths=[60, 150, 80, 142, 120])
        t.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#00ADB5')),
            ('TEXTCOLOR', (0,0), (-1,0), colors.whitesmoke),
            ('ALIGN', (0,0), (-1,-1), 'CENTER'),
            ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
            ('FONTSIZE', (0,0), (-1,0), 10),
            ('BOTTOMPADDING', (0,0), (-1,0), 8),
            ('TOPPADDING', (0,0), (-1,0), 8),
            ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#DDDDDD')),
            ('FONTNAME', (0,1), (-1,-1), 'Helvetica'),
            ('FONTSIZE', (0,1), (-1,-1), 9),
            ('BOTTOMPADDING', (0,1), (-1,-1), 6),
            ('TOPPADDING', (0,1), (-1,-1), 6),
            ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, colors.HexColor('#F9FAFB')])
        ]))
        
        story.append(t)
        doc.build(story)
        logger.info("Successfully generated PDF report file.")
        return True
    except Exception as e:
        logger.error("Failed to generate PDF report: %s", e)
        return False
